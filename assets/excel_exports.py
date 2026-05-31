"""Xuất Excel thiết bị & tiêu thụ."""

from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook

from core.excel_reports import (
    apply_table_column_widths,
    no_wrap_col_indices,
    write_admin_header,
    write_footer_signature,
    write_section_title,
    write_summary_block,
    write_table_header,
    write_table_row,
    workbook_to_response,
)

DEVICE_TYPE_LABELS = {
    'TRANSFORMER': 'Trạm biến áp',
    'DISTRIBUTION_BOX': 'Tủ điện / tủ phân phối',
    'ELECTRIC_POLE': 'Trụ điện',
    'ELECTRIC_JUNCTION': 'Điểm nối điện',
    'ELECTRIC_METER': 'Công tơ điện',
    'WATER_TANK': 'Bể nước',
    'PUMP_STATION': 'Trạm bơm',
    'MAIN_VALVE': 'Van tổng',
    'BRANCH_VALVE': 'Van nhánh',
    'WATER_JUNCTION': 'Điểm nối nước',
    'WATER_METER': 'Đồng hồ nước',
    'VALVE': 'Van nước (cũ)',
}

STATUS_LABELS = {
    'ACTIVE': 'Hoạt động',
    'FAULT': 'Lỗi trực tiếp',
    'MAINTENANCE': 'Đang bảo trì',
    'INACTIVE': 'Ngưng hoạt động',
}

# Thứ tự sheet báo cáo thiết bị theo quận/huyện Đà Nẵng
DEVICE_EXPORT_DISTRICTS = [
    'Thanh Khê',
    'Cẩm Lệ',
    'Sơn Trà',
    'Liên Chiểu',
    'Ngũ Hành Sơn',
    'Hòa Vang',
    'Hải Châu',
]

DEVICE_TABLE_HEADERS = [
    'STT', 'Mã thiết bị', 'Tên thiết bị', 'Loại hạ tầng', 'Phường/Xã', 'Quận/Huyện',
    'Địa chỉ', 'Vĩ độ', 'Kinh độ', 'Trạng thái', 'Hoạt động', 'Cập nhật lần cuối',
]


def _device_district(device) -> str | None:
    if device.ward_id:
        return device.ward.district
    return None


def _sheet_title_for_district(district: str) -> str:
    """Tên sheet Excel (tối đa 31 ký tự)."""
    return district[:31]


def _write_device_district_sheet(ws, devices, district: str, user) -> None:
    """Ghi một sheet báo cáo thiết bị cho một quận/huyện."""
    headers = DEVICE_TABLE_HEADERS
    num_cols = len(headers)
    device_list = list(devices)

    active = sum(1 for d in device_list if d.status == 'ACTIVE')
    fault = sum(1 for d in device_list if d.status == 'FAULT')
    maint = sum(1 for d in device_list if d.status == 'MAINTENANCE')

    row = write_admin_header(
        ws,
        report_title=f'Báo cáo danh mục thiết bị hạ tầng điện — nước',
        report_subtitle=f'(Quận/Huyện {district} — bảng kê chi tiết theo phường/xã)',
        report_number=f'HTĐN/TB/{datetime.now():%m/%Y}',
        num_cols=num_cols,
    )

    write_section_title(ws, row, 1, num_cols, 'I. THÔNG TIN TỔNG HỢP')
    row += 1
    row = write_summary_block(ws, row, 1, num_cols, [
        ('Quận/Huyện', district),
        ('Tổng số thiết bị', str(len(device_list))),
        ('Đang hoạt động', str(active)),
        ('Đang lỗi / sự cố', str(fault)),
        ('Đang bảo trì', str(maint)),
        ('Thời điểm xuất báo cáo', timezone.localtime().strftime('%d/%m/%Y %H:%M')),
        ('Người xuất báo cáo', getattr(user, 'username', '') or '—'),
    ])
    row += 1

    write_section_title(ws, row, 1, num_cols, 'II. BẢNG KÊ CHI TIẾT')
    row += 1
    table_header_row = row
    row = write_table_header(ws, row, headers)

    for idx, d in enumerate(device_list, start=1):
        updated = timezone.localtime(d.updated_at).strftime('%d/%m/%Y %H:%M') if d.updated_at else '—'
        row = write_table_row(ws, row, [
            idx,
            d.code,
            d.name,
            DEVICE_TYPE_LABELS.get(d.device_type, d.device_type),
            d.ward.name if d.ward_id else (d.area or '—'),
            d.ward.district if d.ward_id else '—',
            d.address or '—',
            round(d.latitude, 6),
            round(d.longitude, 6),
            STATUS_LABELS.get(d.status, d.status),
            'Có' if d.is_active else 'Không',
            updated,
        ], center_cols={0, 7, 8, 10}, no_wrap_cols=no_wrap_col_indices(headers))

    table_last_row = row - 1 if device_list else table_header_row
    write_footer_signature(ws, row + 2, num_cols, prepared_by=getattr(user, 'username', ''))
    apply_table_column_widths(
        ws, headers=headers, table_header_row=table_header_row, table_last_row=table_last_row,
    )


def export_devices_excel(devices, user) -> 'HttpResponse':
    """Xuất workbook 7 sheet — mỗi quận/huyện một sheet."""
    device_list = list(devices)

    by_district: dict[str, list] = {name: [] for name in DEVICE_EXPORT_DISTRICTS}
    for device in device_list:
        district = _device_district(device)
        if district in by_district:
            by_district[district].append(device)
        # Thiết bị chưa gán quận: bỏ qua hoặc có thể bổ sung sheet riêng sau

    wb = Workbook()
    wb.remove(wb.active)

    for district in DEVICE_EXPORT_DISTRICTS:
        ws = wb.create_sheet(title=_sheet_title_for_district(district))
        district_devices = sorted(by_district[district], key=lambda d: (d.ward.name if d.ward_id else '', d.code))
        _write_device_district_sheet(ws, district_devices, district, user)

    filename = f'bao_cao_thiet_bi_{datetime.now():%Y%m%d_%H%M}.xlsx'
    return workbook_to_response(wb, filename)


def _consumption_filter_desc(filters: dict | None) -> list[str]:
    filter_desc = []
    if not filters:
        return filter_desc
    if filters.get('device_code'):
        filter_desc.append(f"Mã thiết bị: {filters['device_code']}")
    if filters.get('device_type'):
        filter_desc.append(f"Loại: {filters['device_type']}")
    if filters.get('start_date'):
        filter_desc.append(f"Từ tháng: {filters['start_date']}")
    if filters.get('end_date'):
        filter_desc.append(f"Đến tháng: {filters['end_date']}")
    return filter_desc


def _write_consumption_detail_sheet(
    ws,
    log_list,
    *,
    user,
    report_title: str,
    report_subtitle: str,
    unit: str,
    filters: dict | None = None,
) -> None:
    headers = [
        'STT', 'Mã thiết bị', 'Tên thiết bị', 'Loại', 'Phường/Xã', 'Quận/Huyện',
        'Tháng ghi nhận', 'Chỉ số tiêu thụ', 'Đơn vị', 'Ngày nhập liệu',
    ]
    num_cols = len(headers)
    filter_desc = _consumption_filter_desc(filters)
    total_value = sum(log.value for log in log_list)

    row = write_admin_header(
        ws,
        report_title=report_title,
        report_subtitle=report_subtitle,
        report_number=f'HTĐN/TT/{datetime.now():%m/%Y}',
        num_cols=num_cols,
    )

    write_section_title(ws, row, 1, num_cols, 'I. THÔNG TIN TỔNG HỢP')
    row += 1
    summary = [
        ('Tổng số bản ghi', str(len(log_list))),
        ('Tổng chỉ số tiêu thụ', f'{total_value:,.2f} {unit}'),
        ('Thời điểm xuất báo cáo', timezone.localtime().strftime('%d/%m/%Y %H:%M')),
        ('Người xuất báo cáo', getattr(user, 'username', '') or '—'),
    ]
    if filter_desc:
        summary.insert(0, ('Điều kiện lọc', '; '.join(filter_desc)))
    row = write_summary_block(ws, row, 1, num_cols, summary)
    row += 1

    write_section_title(ws, row, 1, num_cols, 'II. CHI TIẾT CHỈ SỐ TIÊU THỤ')
    row += 1
    table_header_row = row
    row = write_table_header(ws, row, headers)

    for idx, log in enumerate(log_list, start=1):
        dev = log.device
        date_str = log.date.strftime('%m/%Y') if log.date else '—'
        created = timezone.localtime(log.created_at).strftime('%d/%m/%Y') if log.created_at else '—'
        row = write_table_row(ws, row, [
            idx,
            dev.code,
            dev.name,
            DEVICE_TYPE_LABELS.get(dev.device_type, dev.device_type),
            dev.ward.name if dev.ward_id else '—',
            dev.ward.district if dev.ward_id else '—',
            date_str,
            round(log.value, 2),
            unit,
            created,
        ], center_cols={0, 6, 7}, no_wrap_cols=no_wrap_col_indices(headers))

    table_last_row = row - 1 if log_list else table_header_row
    write_footer_signature(ws, row + 2, num_cols, prepared_by=getattr(user, 'username', ''))
    apply_table_column_widths(
        ws, headers=headers, table_header_row=table_header_row, table_last_row=table_last_row,
    )


def export_consumption_excel(logs, user, *, filters: dict | None = None) -> 'HttpResponse':
    log_list = list(logs.select_related('device', 'device__ward'))
    electric_logs = [l for l in log_list if l.device.device_type == 'ELECTRIC_METER']
    water_logs = [l for l in log_list if l.device.device_type == 'WATER_METER']

    wb = Workbook()
    wb.remove(wb.active)

    ws_elec = wb.create_sheet(title='Tieu thu dien')
    _write_consumption_detail_sheet(
        ws_elec,
        electric_logs,
        user=user,
        report_title='Báo cáo tiêu thụ điện',
        report_subtitle='(Chỉ số công tơ điện — đơn vị kWh)',
        unit='kWh',
        filters=filters,
    )

    ws_water = wb.create_sheet(title='Tieu thu nuoc')
    _write_consumption_detail_sheet(
        ws_water,
        water_logs,
        user=user,
        report_title='Báo cáo tiêu thụ nước',
        report_subtitle='(Chỉ số đồng hồ nước — đơn vị m³)',
        unit='m³',
        filters=filters,
    )

    filename = f'bao_cao_tieu_thu_{datetime.now():%Y%m%d_%H%M}.xlsx'
    return workbook_to_response(wb, filename)
