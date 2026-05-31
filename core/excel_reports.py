"""Xuất Excel theo mẫu báo cáo hành chính Việt Nam."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORG_LINE = 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM'
MOTTO_LINE = 'Độc lập - Tự do - Hạnh phúc'
UNIT_NAME = 'Ban Quản lý Hạ tầng Điện — Nước thành phố Đà Nẵng'
CITY_DEFAULT = 'Đà Nẵng, ngày {day} tháng {month} năm {year}'

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT = Font(name='Times New Roman', bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(name='Times New Roman', bold=True, size=14)
SECTION_FONT = Font(name='Times New Roman', bold=True, size=12)
BODY_FONT = Font(name='Times New Roman', size=11)
NARROW_FONT = Font(name='Times New Roman', size=10)
THIN_BORDER = Border(
    left=Side(style='thin', color='666666'),
    right=Side(style='thin', color='666666'),
    top=Side(style='thin', color='666666'),
    bottom=Side(style='thin', color='666666'),
)


def _merge_write(ws, row, col_start, col_end, value, font=None, align=None):
    if col_end > col_start:
        ws.merge_cells(
            start_row=row, start_column=col_start,
            end_row=row, end_column=col_end,
        )
    cell = ws.cell(row=row, column=col_start, value=value)
    cell.font = font or BODY_FONT
    cell.alignment = align or Alignment(horizontal='center', vertical='center', wrap_text=True)
    return cell


def write_admin_header(
    ws,
    *,
    report_title: str,
    report_subtitle: str = '',
    recipient: str = 'Ban Giám đốc / Lãnh đạo đơn vị',
    report_number: str = '',
    num_cols: int = 10,
):
    """Phần đầu trang kiểu văn bản hành chính."""
    _merge_write(
        ws, 1, 1, num_cols, ORG_LINE,
        Font(name='Times New Roman', bold=True, size=11),
        Alignment(horizontal='center', vertical='center'),
    )
    _merge_write(
        ws, 2, 1, num_cols, MOTTO_LINE,
        Font(name='Times New Roman', bold=True, size=11, underline='single'),
        Alignment(horizontal='center', vertical='center'),
    )

    unit = ws.cell(row=4, column=1, value=UNIT_NAME)
    unit.font = Font(name='Times New Roman', bold=True, size=11)
    unit.alignment = Alignment(horizontal='left', vertical='center')

    if report_number:
        num_cell = ws.cell(row=5, column=1, value=f'Số: {report_number}')
        num_cell.font = BODY_FONT
        num_cell.alignment = Alignment(horizontal='left')

    title_row = 7
    _merge_write(ws, title_row, 1, num_cols, report_title.upper(), TITLE_FONT)

    if report_subtitle:
        _merge_write(
            ws, title_row + 1, 1, num_cols, report_subtitle,
            Font(name='Times New Roman', italic=True, size=11),
        )

    _merge_write(
        ws, title_row + 2, 1, num_cols, f'(Kính gửi: {recipient})',
        Font(name='Times New Roman', italic=True, size=11),
    )

    return title_row + 4


def write_section_title(ws, row, col_start, col_end, title: str):
    _merge_write(ws, row, col_start, col_end, title, SECTION_FONT, Alignment(horizontal='left', vertical='center'))


def write_summary_block(ws, start_row, col_start, col_end, items: list[tuple[str, str]]):
    """items: [(label, value), ...]"""
    row = start_row
    label_end = col_start + 1
    value_start = label_end + 1
    for label, value in items:
        ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=label_end)
        lbl = ws.cell(row=row, column=col_start, value=label)
        lbl.font = Font(name='Times New Roman', bold=True, size=11)
        lbl.alignment = Alignment(horizontal='left', vertical='center')
        lbl.border = THIN_BORDER

        ws.merge_cells(start_row=row, start_column=value_start, end_row=row, end_column=col_end)
        val = ws.cell(row=row, column=value_start, value=value)
        val.font = BODY_FONT
        val.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        val.border = THIN_BORDER
        row += 1
    return row


# Cột không xuống dòng — căn theo độ dài nội dung thực tế
NO_WRAP_HEADERS = frozenset({'STT', 'Mã thiết bị', 'Mã SC'})


def no_wrap_col_indices(headers: list[str]) -> set[int]:
    return {i for i, h in enumerate(headers) if h in NO_WRAP_HEADERS}


def write_table_header(ws, row, headers: list[str], start_col: int = 1):
    for idx, text in enumerate(headers):
        col = start_col + idx
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=text not in NO_WRAP_HEADERS,
        )
        cell.border = THIN_BORDER
    return row + 1


def write_table_row(ws, row, values: list, start_col: int = 1, center_cols: set[int] | None = None, no_wrap_cols: set[int] | None = None):
    center_cols = center_cols or set()
    if no_wrap_cols is None:
        no_wrap_cols = {0}
    for idx, value in enumerate(values):
        col = start_col + idx
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = NARROW_FONT
        cell.border = THIN_BORDER
        h = 'center' if idx in center_cols else 'left'
        cell.alignment = Alignment(
            horizontal=h,
            vertical='center',
            wrap_text=idx not in no_wrap_cols,
        )
    return row + 1


# Độ rộng cố định (đơn vị ký tự Excel) cho các cột ngắn / thường gặp
TABLE_COLUMN_PRESETS: dict[str, float] = {
    'STT': 5,
    'Mã SC': 8,
    'Mã thiết bị': 18,
    'Hoạt động': 10,
    'Đơn vị': 8,
    'Vĩ độ': 12,
    'Kinh độ': 12,
    'Loại': 10,
    'Mức độ': 12,
    'Tháng ghi nhận': 13,
    'Chỉ số tiêu thụ': 14,
    'Ngày nhập liệu': 14,
    'Ngày phát sinh': 17,
    'Ngày xử lý xong': 17,
    'Quận/Huyện': 14,
    'Phường/Xã': 16,
    'Trạng thái': 15,
    'Người báo cáo': 14,
    'KTV phụ trách': 14,
    'Cập nhật lần cuối': 17,
}


def apply_table_column_widths(
    ws,
    *,
    headers: list[str],
    table_header_row: int,
    table_last_row: int | None = None,
    min_width: float = 8,
    max_width: float = 34,
):
    """Căn độ rộng cột chỉ theo vùng bảng (tránh STT bị giãn bởi header văn bản phía trên)."""
    last_row = table_last_row if table_last_row is not None else ws.max_row
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        preset = TABLE_COLUMN_PRESETS.get(header)

        if header in NO_WRAP_HEADERS and header != 'STT':
            # Mã thiết bị / Mã SC: rộng vừa đủ nội dung, luôn một dòng
            floor = preset or min_width
            max_len = len(header) + 1
            for row in range(table_header_row + 1, last_row + 1):
                val = ws.cell(row=row, column=idx).value
                if val is None:
                    continue
                max_len = max(max_len, len(str(val).split('\n', 1)[0]) + 1)
            ws.column_dimensions[letter].width = max(floor, min(max_len, max_width))
            continue

        if preset is not None:
            ws.column_dimensions[letter].width = preset
            continue
        max_len = min(len(header) + 2, max_width)
        for row in range(table_header_row, last_row + 1):
            val = ws.cell(row=row, column=idx).value
            if val is None:
                continue
            line = str(val).split('\n', 1)[0]
            max_len = max(max_len, min(len(line) + 2, max_width))
        ws.column_dimensions[letter].width = max(min_width, min(max_len, max_width))


def autosize_columns(ws, min_width=8, max_width=42):
    """Giữ tương thích — ưu tiên dùng apply_table_column_widths cho báo cáo."""
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = min_width
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col_idx).value
            if val is None:
                continue
            max_len = max(max_len, min(len(str(val)) + 2, max_width))
        ws.column_dimensions[letter].width = max_len


def write_footer_signature(ws, row, col_end, *, prepared_by: str = ''):
    sign_col = max(col_end - 2, 1)
    _merge_write(
        ws, row, sign_col, col_end,
        CITY_DEFAULT.format(day=datetime.now().day, month=datetime.now().month, year=datetime.now().year),
        Font(name='Times New Roman', italic=True, size=11),
    )
    _merge_write(ws, row + 1, sign_col, col_end, 'NGƯỜI LẬP BIỂU', Font(name='Times New Roman', bold=True, size=11))
    if prepared_by:
        _merge_write(
            ws, row + 5, sign_col, col_end, prepared_by,
            Font(name='Times New Roman', bold=True, size=11),
        )


def workbook_to_response(wb: Workbook, filename: str) -> HttpResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
